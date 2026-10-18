#!/usr/bin/env python3
"""CIRUS Origin: generates the downloadable templates in site/sablony/.

  cirus-origin-vstupni-sablona.xlsx   input template (sheet "Data" + sheet "Návod")
  cirus-origin-vstupni-sablona.csv    the same columns as CSV (semicolon, UTF-8 with BOM)
  prohlaseni-o-puvodu-cs.docx  supplier's declaration of non-preferential origin, Czech
  prohlaseni-o-puvodu-sk.docx  the same, Slovak
  prohlaseni-o-puvodu-en.docx  the same, English
  README.txt                   what is in the folder

Needs Python 3 with openpyxl and python-docx. Run: python3 tools/make-templates.py
The brand name and web address are read from config.json, so run it again after changing them.
The output is deterministic (fixed document dates), so re-running does not create spurious diffs.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
import re
import sys
import zipfile
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover
    sys.exit("Chybí openpyxl: pip install openpyxl")

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:  # pragma: no cover
    sys.exit("Chybí python-docx: pip install python-docx")

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "site" / "sablony"
FIXED_TIME = dt.datetime(2026, 10, 18, 12, 0, 0)

# design tokens (site/assets/css/tokens.css, light mode)
INK = "0B1220"
MUTED = "5B6472"
LINE = "E4E7EB"
SUNK = "F4F5F7"
ACCENT = "0F3D73"
ACCENT_SOFT = "EAF0F7"
MARK = "0F3D73"
MARK_SOFT = "EEF0F3"


def load_config() -> dict:
    try:
        return json.loads((ROOT / "config.json").read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


CFG = load_config()
BRAND = str(CFG.get("BRAND") or "CIRUS Origin")
WEB = str(CFG.get("WEB") or "").strip().removeprefix("https://").removeprefix("http://").rstrip("/")
WEB_OK = bool(WEB) and "[" not in WEB and " " not in WEB and "." in WEB
CALC_URL = f"https://{WEB}/kalkulacka/" if WEB_OK else "kalkulacka/index.html"

# ---------------------------------------------------------------- input template

# Column headers recognized by calc-core.js detectColumns(); keep in sync with docs/SPEC.md section 7.
COLUMNS = [
    # header, width, field, required, help cs, help en
    ("Číslo dílu", 16, "partNo", "ne",
     "Vaše číslo dílu nebo pseudonym. Nemusí být skutečné.",
     "Your part number or a pseudonym."),
    ("Popis", 38, "description", "ne",
     "Obecný popis dílu. Nevkládejte technické specifikace.",
     "Generic description. No technical specifications."),
    ("Dodavatel", 28, "supplier", "doporučeno",
     "Přímý dodavatel (tier-1), od kterého díl kupujete. U vlastní výroby napište „vlastní výroba“.",
     "The direct (tier-1) supplier you buy from. For in-house production write \"in-house\"."),
    ("Země dodavatele", 12, "supplierCountry", "doporučeno",
     "Kód ISO sídla dodavatele (CZ, DE, US…). Použije se jako odhad původu, když původ chybí.",
     "ISO code of the supplier's registered office. Used as an estimate when origin is missing."),
    ("Kategorie", 26, "category", "doporučeno",
     "Jedna z pěti kategorií Komise (výběr ze seznamu).",
     "One of the Commission's five categories (drop-down)."),
    ("Množství", 10, "qty", "ano",
     "Počet kusů na jeden výrobek. Prázdné = 1.",
     "Quantity per end product. Empty = 1."),
    ("Jednotková cena", 15, "unitCost", "ano",
     "Cena za kus podle faktury, smlouvy nebo celního dokladu. Bez dopravy, cla a DPH.",
     "Price per unit from the invoice, contract or customs document. Without transport, duty and VAT."),
    ("Měna", 8, "currency", "ne",
     "Kód měny ISO (CZK, EUR, USD…). Prázdné = měna projektu.",
     "ISO currency code. Empty = project currency."),
    ("Země původu", 12, "origin", "ano, pokud ji znáte",
     "Nepreferenční původ dílu, kód ISO. Pozor: není to sídlo dodavatele.",
     "Non-preferential origin of the part, ISO code. Not the supplier's country."),
    ("Doklad", 16, "evidence", "ano, pokud je",
     "Čím je původ doložen (výběr ze seznamu).",
     "What documents the origin (drop-down)."),
    ("Platnost do", 13, "validUntil", "u prohlášení",
     "Konec platnosti prohlášení dodavatele (RRRR-MM-DD).",
     "Expiry of the supplier's declaration (YYYY-MM-DD)."),
    ("Odhad původu", 12, "expectedOrigin", "ne",
     "Váš nejlepší odhad původu, dokud ho nemáte doložený. Kód ISO.",
     "Your best estimate of origin until it is documented. ISO code."),
    ("Typ", 15, "type", "ano",
     "díl = nakupovaný díl (tier-1); vlastní výroba = podsestava vyrobená ve vlastní režii; surovina = nezapočítává se.",
     "part = purchased tier-1 part; in-house = sub-assembly made in-house; raw material = excluded."),
    ("Poznámka", 36, "note", "ne",
     "Cokoli, co pomůže při ověření.",
     "Anything that helps verification."),
]
HEADERS = [c[0] for c in COLUMNS]

CATEGORY_VALUES = [
    "konstrukce a mechanika",
    "pohon a energie",
    "přístroje a monitorování",
    "bezpečnost a ochrana",
    "elektronika a komunikace",
]
EVIDENCE_VALUES = ["celní doklad", "prohlášení", "smlouva", "odhad", "žádný"]
TYPE_VALUES = ["díl", "vlastní výroba", "surovina"]

EXAMPLE_NOTE = "Příklad, řádek smažte"
EXAMPLES = [
    ["ABC-1001", "Svařovaný rám", "Strojírna Vzor s.r.o.", "CZ", "konstrukce a mechanika", 1, 185000, "CZK",
     "CZ", "prohlášení", dt.date(2027, 12, 31), "", "díl", EXAMPLE_NOTE],
    ["ABC-2040", "Řídicí jednotka", "Vzorový obchodník GmbH", "DE", "elektronika a komunikace", 2, 3400.5, "EUR",
     "", "žádný", None, "JP", "díl", f"{EXAMPLE_NOTE}. Obchodník, výrobce zatím neznámý."],
    ["ABC-9001", "Montáž podsestavy rámu (práce a režie)", "vlastní výroba", "CZ", "konstrukce a mechanika", 1,
     64000, "CZK", "CZ", "", None, "", "vlastní výroba",
     f"{EXAMPLE_NOTE}. Bez nakoupených dílů, které už jsou v seznamu."],
]
DATA_ROWS = 1000  # rows covered by validation and formatting


def fmt_csv_value(v) -> str:
    if v is None:
        return ""
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float):
        return f"{v:.2f}".replace(".", ",")
    return str(v)


def thin(color: str = LINE) -> Border:
    s = Side(style="thin", color=color)
    return Border(left=s, right=s, top=s, bottom=s)


def build_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.sheet_properties.tabColor = ACCENT

    head_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    head_fill = PatternFill("solid", fgColor=ACCENT)
    body_font = Font(name="Arial", size=10, color=INK)
    example_font = Font(name="Arial", size=10, color=MUTED, italic=True)

    ws.append(HEADERS)
    for col, (_, width, *_rest) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = thin(ACCENT)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 30

    for r, row in enumerate(EXAMPLES, start=2):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value if value != "" else None)
            cell.font = example_font
    col_of = {c[2]: i for i, c in enumerate(COLUMNS, start=1)}
    for r in range(2, DATA_ROWS + 2):
        ws.cell(row=r, column=col_of["qty"]).number_format = "#,##0.###"
        ws.cell(row=r, column=col_of["unitCost"]).number_format = "#,##0.00"
        ws.cell(row=r, column=col_of["validUntil"]).number_format = "yyyy-mm-dd"
        for key in ("supplierCountry", "origin", "expectedOrigin", "currency"):
            ws.cell(row=r, column=col_of[key]).alignment = Alignment(horizontal="center")
        if r > len(EXAMPLES) + 1:
            for c in range(1, len(COLUMNS) + 1):
                ws.cell(row=r, column=c).font = body_font

    ws.freeze_panes = "A2"
    last = get_column_letter(len(COLUMNS))
    ws.auto_filter.ref = f"A1:{last}{DATA_ROWS + 1}"

    def rng(field: str) -> str:
        letter = get_column_letter(col_of[field])
        return f"{letter}2:{letter}{DATA_ROWS + 1}"

    def list_validation(field: str, values: list[str], title: str, prompt: str) -> None:
        dv = DataValidation(
            type="list",
            formula1='"' + ",".join(values) + '"',
            allow_blank=True,
            showDropDown=False,  # False = show the in-cell arrow (openpyxl naming is inverted)
            showErrorMessage=True,
            errorStyle="warning",
            errorTitle="Hodnota mimo seznam",
            error="Hodnota není v seznamu. Kalkulačka zkusí rozpoznat i jiné výrazy (také anglické). Pokračovat?",
            showInputMessage=True,
            promptTitle=title[:32],
            prompt=prompt[:255],
        )
        dv.add(rng(field))
        ws.add_data_validation(dv)

    list_validation("category", CATEGORY_VALUES, "Kategorie", "Jedna z pěti kategorií podle šablony Komise.")
    list_validation("evidence", EVIDENCE_VALUES, "Doklad o původu",
                    "celní doklad, prohlášení dodavatele, smlouva, odhad, nebo žádný.")
    list_validation("type", TYPE_VALUES, "Typ řádku",
                    "díl = nakupovaný (tier-1); vlastní výroba = práce a režie podsestav; surovina = nezapočítává se.")

    def hint(field: str, title: str, prompt: str) -> None:
        dv = DataValidation(type=None, allow_blank=True, showInputMessage=True, promptTitle=title[:32], prompt=prompt[:255])
        dv.add(rng(field))
        ws.add_data_validation(dv)

    hint("origin", "Země původu", "Nepreferenční původ dílu, kód ISO (CZ, DE, US…). Není to sídlo dodavatele.")
    hint("expectedOrigin", "Odhad původu", "Jen pokud původ zatím nemáte doložený. Kód ISO.")
    hint("supplierCountry", "Země dodavatele", "Kód ISO sídla dodavatele (CZ, DE, US…).")
    hint("unitCost", "Jednotková cena", "Bez dopravy, cla a DPH. Podle faktury, smlouvy nebo celního dokladu.")
    hint("qty", "Množství", "Počet kusů na jeden výrobek.")
    hint("validUntil", "Platnost do", "Konec platnosti prohlášení, RRRR-MM-DD (nejdéle 24 měsíců).")

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"

    build_guide(wb.create_sheet("Návod"))

    wb.properties.creator = BRAND
    wb.properties.lastModifiedBy = BRAND
    wb.properties.title = f"{BRAND}: vstupní šablona"
    wb.properties.created = FIXED_TIME
    wb.properties.modified = FIXED_TIME
    wb.active = 0
    buf = io.BytesIO()
    wb.save(buf)
    write_normalized_zip(path, buf.getvalue())


def lines_height(text: str, per_line: int = 115) -> float | None:
    """Row height for wrapped text in merged cells (Excel does not auto-fit those)."""
    n = -(-len(text) // per_line)
    return None if n <= 1 else 14.5 * n


def build_guide(ws) -> None:
    ws.sheet_properties.tabColor = MARK
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 90
    ws.sheet_view.showGridLines = False

    title = Font(name="Arial", size=16, bold=True, color=INK)
    h2 = Font(name="Arial", size=12, bold=True, color=ACCENT)
    body = Font(name="Arial", size=10, color=INK)
    small = Font(name="Arial", size=9, color=MUTED, italic=True)
    th = Font(name="Arial", size=10, bold=True, color=INK)
    wrap = Alignment(wrap_text=True, vertical="top")

    row = 1

    def put(text: str, font=body, merge: bool = True, height: float | None = None) -> None:
        nonlocal row
        ws.cell(row=row, column=1, value=text).font = font
        ws.cell(row=row, column=1).alignment = wrap
        if merge:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        if height:
            ws.row_dimensions[row].height = height
        row += 1

    def table(head: tuple[str, str, str], rows: list[tuple[str, str, str]]) -> None:
        nonlocal row
        for c, v in enumerate(head, start=1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.font = th
            cell.fill = PatternFill("solid", fgColor=SUNK)
            cell.border = thin()
            cell.alignment = wrap
        row += 1
        for r in rows:
            for c, v in enumerate(r, start=1):
                cell = ws.cell(row=row, column=c, value=v)
                cell.font = Font(name="Arial", size=10, color=INK, bold=(c == 1))
                cell.border = thin()
                cell.alignment = wrap
            ws.row_dimensions[row].height = 15 * max(1, -(-len(r[2]) // 95))
            row += 1

    put(f"{BRAND}: vstupní šablona pro výpočet podílu komponent mimo EU", title, height=24)
    put("Pravidlo 35 % podle nařízení SAFE a EDIP. Data vyplňte na listu „Data“.", small)
    row += 1

    put("Jak vyplnit", h2)
    for line in [
        "1. Jeden řádek = jeden díl, který vaše firma nakupuje přímo (tier-1), nebo jedna podsestava z vlastní výroby.",
        "2. Uveďte všechny přímé nákupy bez ohledu na to, jak je díl důležitý. Díly vašich dodavatelů (tier-2) neuvádějte.",
        "3. Cena za kus podle faktury, smlouvy nebo celního dokladu, bez dopravy, cla a DPH. Množství na jeden výrobek.",
        "4. Země původu je nepreferenční původ dílu (čl. 60 celního kodexu Unie), ne sídlo dodavatele. U dílů koupených "
        "od obchodníka rozhoduje, kde díl vyrobil výrobce.",
        "5. Doklad: celní doklad (dovoz), prohlášení dodavatele, smlouva, odhad, nebo žádný. U prohlášení vyplňte i „Platnost do“.",
        "6. Když původ zatím nemáte doložený, vyplňte „Odhad původu“. Kalkulačka spočítá nejhorší, očekávaný i nejlepší případ.",
        "7. Vlastní výroba: celkové výrobní náklady podsestav vyrobených ve vlastní režii (práce a režie), bez nakoupených "
        "dílů, které už jsou v seznamu, a bez konečné montáže. Používá se jen pro metodu (b).",
        "8. Suroviny nejsou komponenty (EDIP, recitál 27). Uveďte je s typem „surovina“, nebo je vynechte.",
        "9. První tři řádky jsou příklad. Před výpočtem je smažte.",
    ]:
        put(line, height=lines_height(line))
    row += 1

    put("Sloupce", h2)
    table(("Sloupec", "Povinný", "Co vyplnit"), [(c[0], c[3], c[4]) for c in COLUMNS])
    row += 1

    put("Hodnoty ze seznamů", h2)
    table(("Sloupec", "Hodnota", "Význam"), [
        ("Kategorie", "konstrukce a mechanika", "structural and mechanical"),
        ("", "pohon a energie", "power and energy"),
        ("", "přístroje a monitorování", "instrumentation and monitoring"),
        ("", "bezpečnost a ochrana", "safety and protection"),
        ("", "elektronika a komunikace", "electronics and comms"),
        ("Doklad", "celní doklad", "původ dovezeného dílu podle celního dokladu"),
        ("", "prohlášení", "prohlášení dodavatele o nepreferenčním původu (vzor je ke stažení na webu)"),
        ("", "smlouva", "původ uvedený v kupní smlouvě"),
        ("", "odhad", "původ je uvedený, ale ničím nedoložený"),
        ("", "žádný", "doklad chybí"),
        ("Typ", "díl", "nakupovaný díl (tier-1), započítává se"),
        ("", "vlastní výroba", "podsestava vyrobená ve vlastní režii, jen pro metodu (b)"),
        ("", "surovina", "nezapočítává se"),
    ])
    row += 1

    put("Co dál", h2)
    put(f"Otevřete kalkulačku {CALC_URL}, klikněte na „Nahrát CSV nebo XLSX“ a vyberte tento soubor. "
        "Nebo v Excelu označte data včetně řádku se záhlavím, zkopírujte je a vložte do kalkulačky.", height=29)
    put("Kalkulačka běží ve vašem prohlížeči a data nikam neodesílá. Šablona neobsahuje makra.")
    put("Pro výpočet stačí obchodní údaje. Nevkládejte výkresy, technické specifikace ani utajované informace.")
    put("Výsledek je výpočet z vašich dat a dokladů, ne právní stanovisko.", small)
    row += 1

    put("English", h2)
    for line in [
        "Input template for calculating the share of components originating outside the EU (35 % rule under SAFE and EDIP).",
        "One row = one part your company buys directly (tier-1), or one sub-assembly made in-house. "
        "List all direct purchases regardless of importance; do not list your suppliers' parts (tier-2).",
        "Unit price from the invoice, contract or customs document, without transport, duty and VAT. Quantity per end product.",
        "Country of origin = non-preferential origin of the part (Art. 60 Union Customs Code), not the supplier's country.",
        "Evidence: customs document, declaration, contract, estimate, or none. For declarations fill in the expiry date.",
        "Type: part (tier-1, counted), in-house (sub-assembly labour and overhead, method b only), raw material (excluded).",
        "The headers may stay in Czech; the calculator also recognizes English headers and values. "
        "The first three rows are examples, delete them before calculating.",
        f"Open the calculator at {CALC_URL} and upload this file. It runs in your browser and sends no data anywhere.",
    ]:
        put(line, height=lines_height(line))

    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def build_csv(path: Path) -> None:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", lineterminator="\r\n", quoting=csv.QUOTE_MINIMAL)
    w.writerow(HEADERS)
    for row in EXAMPLES:
        w.writerow([fmt_csv_value(v) for v in row])
    path.write_bytes(b"\xef\xbb\xbf" + buf.getvalue().encode("utf-8"))


# ------------------------------------------------------------ supplier declaration

# Texts mirror site/assets/js/declaration.js (SPEC section 8), so the Word template and the
# document generated by the calculator say the same thing.
DECL = {
    "cs": {
        "file": "prohlaseni-o-puvodu-cs.docx",
        "title": "Prohlášení dodavatele o nepreferenčním původu zboží",
        "template": "Vzor",
        "supplier": "Dodavatel",
        "buyer": "Odběratel",
        "name": "Obchodní jméno",
        "address": "Sídlo",
        "ids": "IČO / DIČ",
        "contact": "Kontaktní osoba",
        "intro": "Dodavatel prohlašuje, že níže uvedené zboží má nepreferenční původ v uvedené zemi ve smyslu "
                 "čl. 60 nařízení Evropského parlamentu a Rady (EU) č. 952/2013, kterým se stanoví celní kodex Unie. "
                 "U zboží, na jehož výrobě se podílela více než jedna země, je původ určen podle místa poslední "
                 "podstatné, hospodářsky odůvodněné přeměny.",
        "no": "Č.",
        "partNo": "Číslo dílu",
        "description": "Popis",
        "origin": "Země původu",
        "originNote": "(vyplní dodavatel)",
        "basis": "Základ určení původu",
        "basisNote": "(A, B nebo C)",
        "basisLegend": "Základ určení původu:",
        "basisOptions": ["A – vyrobeno u nás", "B – poslední podstatná přeměna u nás",
                         "C – převzato od výrobce – uveďte výrobce"],
        "validity": "Prohlášení platí do:",
        "validityNote": "(nejdéle 24 měsíců ode dne podpisu)",
        "commitment": "Dodavatel se zavazuje bez zbytečného odkladu písemně oznámit odběrateli každou změnu, "
                      "která má vliv na původ uvedeného zboží.",
        "place": "Místo",
        "date": "Datum",
        "signatory": "Jméno a funkce",
        "signature": "Podpis",
        "stamp": "Razítko",
        "footnote": "Vzor. Nejde o preferenční prohlášení podle prováděcího nařízení (EU) 2015/2447.",
        "lang": "cs-CZ",
    },
    "sk": {
        "file": "prohlaseni-o-puvodu-sk.docx",
        "title": "Vyhlásenie dodávateľa o nepreferenčnom pôvode tovaru",
        "template": "Vzor",
        "supplier": "Dodávateľ",
        "buyer": "Odberateľ",
        "name": "Obchodné meno",
        "address": "Sídlo",
        "ids": "IČO / DIČ",
        "contact": "Kontaktná osoba",
        "intro": "Dodávateľ vyhlasuje, že nižšie uvedený tovar má nepreferenčný pôvod v uvedenej krajine v zmysle "
                 "článku 60 nariadenia Európskeho parlamentu a Rady (EÚ) č. 952/2013, ktorým sa ustanovuje Colný "
                 "kódex Únie. Pri tovare, na ktorého výrobe sa podieľala viac ako jedna krajina, je pôvod určený "
                 "podľa miesta posledného podstatného, ekonomicky odôvodneného spracovania alebo opracovania.",
        "no": "Č.",
        "partNo": "Číslo dielu",
        "description": "Popis",
        "origin": "Krajina pôvodu",
        "originNote": "(vyplní dodávateľ)",
        "basis": "Základ určenia pôvodu",
        "basisNote": "(A, B alebo C)",
        "basisLegend": "Základ určenia pôvodu:",
        "basisOptions": ["A – vyrobené u nás", "B – posledné podstatné spracovanie u nás",
                         "C – prevzaté od výrobcu – uveďte výrobcu"],
        "validity": "Vyhlásenie platí do:",
        "validityNote": "(najviac 24 mesiacov odo dňa podpisu)",
        "commitment": "Dodávateľ sa zaväzuje bez zbytočného odkladu písomne oznámiť odberateľovi každú zmenu, "
                      "ktorá má vplyv na pôvod uvedeného tovaru.",
        "place": "Miesto",
        "date": "Dátum",
        "signatory": "Meno a funkcia",
        "signature": "Podpis",
        "stamp": "Pečiatka",
        "footnote": "Vzor. Nejde o preferenčné vyhlásenie podľa vykonávacieho nariadenia (EÚ) 2015/2447.",
        "lang": "sk-SK",
    },
    "en": {
        "file": "prohlaseni-o-puvodu-en.docx",
        "title": "Supplier’s declaration of non-preferential origin of goods",
        "template": "Template",
        "supplier": "Supplier",
        "buyer": "Buyer",
        "name": "Company name",
        "address": "Registered office",
        "ids": "Company ID / VAT number",
        "contact": "Contact person",
        "intro": "The supplier declares that the goods listed below have their non-preferential origin in the "
                 "country stated, within the meaning of Article 60 of Regulation (EU) No 952/2013 of the European "
                 "Parliament and of the Council laying down the Union Customs Code. Where the production of the "
                 "goods involved more than one country, origin is determined by the place of the last substantial, "
                 "economically justified processing or working.",
        "no": "No.",
        "partNo": "Part number",
        "description": "Description",
        "origin": "Country of origin",
        "originNote": "(to be completed by the supplier)",
        "basis": "Basis for origin",
        "basisNote": "(A, B or C)",
        "basisLegend": "Basis for origin:",
        "basisOptions": ["A – manufactured by us", "B – last substantial transformation by us",
                         "C – taken over from the manufacturer – name the manufacturer"],
        "validity": "This declaration is valid until:",
        "validityNote": "(no longer than 24 months from the date of signature)",
        "commitment": "The supplier undertakes to inform the buyer in writing without undue delay of any change "
                      "affecting the origin of the goods listed.",
        "place": "Place",
        "date": "Date",
        "signatory": "Name and position",
        "signature": "Signature",
        "stamp": "Company stamp",
        "footnote": "Template. This is not a preferential origin declaration under Implementing Regulation (EU) 2015/2447.",
        "lang": "en-GB",
    },
}
DECL_ROWS = 10


def shade(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def table_borders(table, color: str = "8A949B", size: int = 4) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def no_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def row_height(row, cm: float) -> None:
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    row.height = Cm(cm)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST


def set_cell(cell, text: str, bold: bool = False, size: float = 9.5, color: str | None = None,
             note: str | None = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if note:
        r2 = p.add_run(f"\n{note}")
        r2.font.size = Pt(8)
        r2.font.color.rgb = RGBColor.from_string(MUTED)


def set_widths(table, widths_cm: list[float]) -> None:
    """Fixed column widths: grid columns (LibreOffice) and cell widths (Word)."""
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    for idx, w in enumerate(widths_cm):
        table.columns[idx].width = Cm(w)
    for row in table.rows:
        for idx, w in enumerate(widths_cm):
            row.cells[idx].width = Cm(w)


def set_language(doc, lang: str) -> None:
    rpr = doc.styles["Normal"].element.get_or_add_rPr()
    el = OxmlElement("w:lang")
    el.set(qn("w:val"), lang)
    el.set(qn("w:eastAsia"), lang)
    rpr.append(el)


def build_docx(lang: str, path: Path) -> None:
    t = DECL[lang]
    doc = Document()

    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    set_language(doc, t["lang"])

    # kicker + title
    p = doc.add_paragraph()
    r = p.add_run(t["template"].upper())
    r.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(MARK)
    p.paragraph_format.space_after = Pt(2)

    p = doc.add_paragraph()
    r = p.add_run(t["title"])
    r.bold = True
    r.font.size = Pt(16)
    p.paragraph_format.space_after = Pt(10)

    # parties
    parties = doc.add_table(rows=5, cols=3)
    parties.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(parties)
    set_cell(parties.cell(0, 0), "")
    set_cell(parties.cell(0, 1), t["supplier"], bold=True)
    set_cell(parties.cell(0, 2), t["buyer"], bold=True)
    for i, key in enumerate(("name", "address", "ids", "contact"), start=1):
        set_cell(parties.cell(i, 0), t[key], size=9, color=MUTED)
        row_height(parties.rows[i], 0.75)
    for c in range(3):
        shade(parties.cell(0, c), SUNK)
    set_widths(parties, [3.6, 6.7, 6.7])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # statement
    p = doc.add_paragraph(t["intro"])
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(8)

    # goods table
    goods = doc.add_table(rows=DECL_ROWS + 1, cols=5)
    goods.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(goods)
    heads = [(t["no"], None), (t["partNo"], None), (t["description"], None),
             (t["origin"], t["originNote"]), (t["basis"], t["basisNote"])]
    for c, (h, note) in enumerate(heads):
        cell = goods.cell(0, c)
        set_cell(cell, h, bold=True, size=9, note=note)
        shade(cell, SUNK)
    repeat_header(goods.rows[0])
    for r_i in range(1, DECL_ROWS + 1):
        set_cell(goods.cell(r_i, 0), str(r_i), size=9, color=MUTED)
        row_height(goods.rows[r_i], 0.75)
        no_split(goods.rows[r_i])
    set_widths(goods, [1.0, 3.4, 5.4, 3.4, 3.8])

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(t["basisLegend"])
    r.bold = True
    r.font.size = Pt(9)
    for opt in t["basisOptions"]:
        r = p.add_run()
        r.add_break()
        r = p.add_run(opt)
        r.font.size = Pt(9)

    # validity + commitment
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    r = p.add_run(f"{t['validity']} ")
    r.bold = True
    p.add_run("______________________ ")
    r = p.add_run(t["validityNote"])
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(MUTED)

    p = doc.add_paragraph(t["commitment"])
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(12)

    # place, date, signature, stamp
    sign = doc.add_table(rows=2, cols=5)
    sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(sign)
    for c, key in enumerate(("place", "date", "signatory", "signature", "stamp")):
        set_cell(sign.cell(0, c), t[key], size=9, color=MARK if key == "stamp" else MUTED)
        shade(sign.cell(0, c), MARK_SOFT if key == "stamp" else SUNK)
    row_height(sign.rows[1], 2.6)
    no_split(sign.rows[0])
    no_split(sign.rows[1])
    set_widths(sign, [3.3, 2.6, 4.1, 3.4, 3.6])

    # footer: footnote + brand
    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.text = ""
    r = fp.add_run(t["footnote"])
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(MUTED)
    if WEB_OK:
        r = fp.add_run(f"  ·  {BRAND}, {WEB}")
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string(MUTED)

    cp = doc.core_properties
    cp.title = t["title"]
    cp.author = BRAND
    cp.last_modified_by = BRAND
    cp.language = t["lang"]
    cp.created = FIXED_TIME
    cp.modified = FIXED_TIME
    cp.revision = 1

    buf = io.BytesIO()
    doc.save(buf)
    write_normalized_zip(path, buf.getvalue())


# ------------------------------------------------------------------- helpers

def write_normalized_zip(path: Path, data: bytes) -> None:
    """Rewrites an OOXML zip with fixed timestamps so that re-running gives identical bytes."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            zi = zipfile.ZipInfo(info.filename, date_time=(2026, 10, 18, 12, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            content = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                content = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:)",
                                 rb"\g<1>2026-10-18T12:00:00Z\g<2>", content)
            dst.writestr(zi, content)
    path.write_bytes(out.getvalue())


# README.txt goes through the site build, which fills in {{BRAND}} and {{WEB}} from config.json.
README = """{{BRAND}}: šablony ke stažení
================================

cirus-origin-vstupni-sablona.xlsx
  Vstupní šablona pro kalkulačku. List „Data“ má sloupce, které kalkulačka
  rozpozná sama, a nabídky hodnot pro kategorii, doklad a typ řádku.
  List „Návod“ popisuje, co do kterého sloupce patří.

cirus-origin-vstupni-sablona.csv
  Stejné sloupce jako CSV (oddělovač středník, UTF-8). Pro systémy, které
  exportují jen CSV.

prohlaseni-o-puvodu-cs.docx, -sk.docx, -en.docx
  Vzor prohlášení dodavatele o nepreferenčním původu zboží (čl. 60 celního
  kodexu Unie) v češtině, slovenštině a angličtině. Vyplníte odběratele
  a díly, dodavatel doplní zemi původu a základ jejího určení, podepíše
  a vrátí. Platnost nejdéle 24 měsíců. Nejde o preferenční prohlášení
  podle prováděcího nařízení (EU) 2015/2447.

Kalkulačka: https://{{WEB}}/kalkulacka/
Kalkulačka běží ve vašem prohlížeči, data se nikam neodesílají.


Templates (English)
-------------------
cirus-origin-vstupni-sablona.xlsx / .csv  input template for the calculator
prohlaseni-o-puvodu-en.docx        supplier's declaration of non-preferential
                                   origin (Art. 60 Union Customs Code)
"""


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    build_xlsx(OUT / "cirus-origin-vstupni-sablona.xlsx")
    build_csv(OUT / "cirus-origin-vstupni-sablona.csv")
    for lang, t in DECL.items():
        build_docx(lang, OUT / t["file"])
    (OUT / "README.txt").write_bytes(README.replace("\n", "\r\n").encode("utf-8"))
    for f in sorted(OUT.iterdir()):
        print(f"  {f.relative_to(ROOT)}  ({f.stat().st_size:,} B)".replace(",", " "))


if __name__ == "__main__":
    main()
